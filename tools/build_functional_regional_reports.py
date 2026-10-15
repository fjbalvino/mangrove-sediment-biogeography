#!/usr/bin/env python3
"""Build separate plain-text Word reports with bold titles and real figures."""
import argparse
import re
from pathlib import Path
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def format_doc(doc,title):
    s=doc.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.left_margin=s.right_margin=Inches(.75)
    s.top_margin=s.bottom_margin=Inches(.7)
    for name in ['Normal','Body Text','Caption','Title','Heading 1','Heading 2','Heading 3']:
        st=doc.styles[name]
        st.font.name='Arial';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
        st.font.italic=False;st.font.underline=False
        st.font.bold=name in ['Title','Heading 1','Heading 2','Heading 3']
        st.paragraph_format.line_spacing=1
        st.paragraph_format.space_before=Pt(0)
        st.paragraph_format.space_after=Pt(6)
        st.paragraph_format.keep_with_next=name in ['Title','Heading 1','Heading 2','Heading 3']
    doc.core_properties.title=title;doc.core_properties.author='MAGENTA'
    doc.core_properties.subject='Figure 5 functional prevalence, selected processes and locality turnover'
    for sec in doc.sections:
        for part in [sec.header,sec.footer]:
            for p in part.paragraphs:p.text=''


def paragraph(doc,text,heading=False):
    text=re.sub(r'\[([^]]+)\]\([^)]+\)',r'\1',text).replace('`','').replace('**','')
    p=doc.add_paragraph(text,style='Heading 1' if heading else 'Normal')
    for r in p.runs:
        r.bold=bool(heading);r.italic=False;r.underline=False
        r.font.name='Arial';r.font.size=Pt(11);r.font.color.rgb=RGBColor(0,0,0)
    return p


def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
    header=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(header)
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for row in t.rows:
        tr=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(tr)
        for c,w in zip(row.cells,widths):
            c.width=Inches(w)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                p.paragraph_format.line_spacing=1
                for r in p.runs:
                    r.font.name='Arial';r.font.size=Pt(11);r.bold=False;r.italic=False;r.underline=False;r.font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph()
    return t


def picture(doc,path,width):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(3)
    r=p.add_run();r.add_picture(str(path),width=Inches(width))
    desc=p._p.xpath('.//wp:docPr')
    if desc:desc[0].set('descr',path.stem.replace('_',' '))


def sections(path):
    result={}
    for part in re.split(r'^## ',path.read_text(),flags=re.M)[1:]:
        head,body=part.split('\n',1)
        result[head]=[x.strip() for x in body.strip().split('\n\n') if x.strip()]
    return result


def page(doc,title):
    doc.add_page_break();paragraph(doc,title,True)


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--results',type=Path,required=True);p.add_argument('--outdir',type=Path,required=True)
    p.add_argument('--main-only',action='store_true',help='Update the main report without rewriting the supplement.')
    a=p.parse_args();a.outdir.mkdir(parents=True,exist_ok=True)
    repo=Path(__file__).resolve().parents[1];tabs=a.results/'tables';plots=a.results/'plots'
    text=(repo/'docs/FUNCTIONAL_REGIONAL_RESULTS.md').read_text()
    title,body=text.split('\n',1);body,caption=body.split('## Figure 5 caption\n')
    doc=Document();format_doc(doc,title.lstrip('# '))
    paragraph(doc,title.lstrip('# '),True)
    for x in body.strip().split('\n\n'):paragraph(doc,x)
    cap=caption.strip().split('\n\n')
    paragraph(doc,cap[0],True)
    for x in cap[1:]:paragraph(doc,x)
    page(doc,'Figure 5')
    picture(doc,plots/'Figure_5_functional_regional_identity_combined.png',7)
    doc.save(a.outdir/'Functional_regional_results.docx')
    if a.main_only:
        print('Replaced main report; figure retained; plain text and bold titles.')
        return

    s=sections(repo/'docs/FUNCTIONAL_REGIONAL_SUPPLEMENTARY_RESULTS.md')
    doc=Document();format_doc(doc,'Supplementary results and methods for Figure 5')
    paragraph(doc,'Supplementary results and methods for Figure 5',True)
    paragraph(doc,'Cohort and persistence reconciliation',True)
    for x in s['Cohort and persistence reconciliation']:paragraph(doc,x)
    paragraph(doc,'Balanced sampling and transformations',True)
    paragraph(doc,'Each of 499 primary selections drew six named localities without replacement in each of Central Indo-Pacific, Tropical Atlantic and Western Indo-Pacific, followed by one randomly selected profile per locality (18 profiles per selection). The alternative used four localities per realm (12 profiles). Identical primary profiles were compared across prevalence groups, persistence intersections and process scores. All 191 profiles were retained when defining KO classes.')
    paragraph(doc,'For each KO fraction separately, counts were CLR-transformed after adding one; Euclidean distances in CLR space were used for Aitchison PERMANOVA. The complementary model conditioned on log₁₀ total KO annotation counts. Its unique realm R² uses total sum of squares as denominator and is not a degrees-of-freedom correction. Total counts can reflect technical and biological differences, so conditioning is a sensitivity analysis. PERMANOVA and spatial-median PERMDISP with small-sample bias adjustment used 499 permutations among selected locality representatives. Repeated selections overlap and do not constitute independent replication.')

    page(doc,'Balanced regional comparisons')
    stats=pd.read_csv(tabs/'180_verified_main_statistics.tsv',sep='\t')
    def statrows(z):
        return [[r.group+f'\n({r.n_KOs:,} KOs)',f'{r.R2_median:.3f}\n({r.R2_q05:.3f}–{r.R2_q95:.3f})',f'{r.adjusted_unique_R2_median:.3f}',f'{r.permanova_P_nominal_count}/499\n({r.permanova_P_nominal_percent:.1f}%)',f'{r.permdisp_P_nominal_count}/499\n({r.permdisp_P_nominal_percent:.1f}%)'] for r in z.itertuples()]
    hdr=['KO group','R² median\n(5th–95th)','Unique realm R²','PERMANOVA\nP ≤ 0.05','PERMDISP\nP ≤ 0.05']
    for design in ['3 x 6','3 x 4']:
        paragraph(doc,('Primary design ' if design=='3 x 6' else 'Alternative design ')+design.replace('x','×'),True)
        z=stats[(stats.comparison=='prevalence')&(stats.design==design)].set_index('group').loc[['Prevalent','Shell','Cloud']].reset_index()
        table(doc,hdr,statrows(z),[1.65,1.7,1.15,1.25,1.25])
    paragraph(doc,'Persistence intersections in the primary design',True)
    z=stats[stats.comparison=='intersection'].set_index('group').loc[['Prevalent','Realm-core Shell','Peripheral Shell','Peripheral Cloud']].reset_index()
    table(doc,hdr,statrows(z),[1.65,1.7,1.15,1.25,1.25])
    paragraph(doc,'Unique realm R² is conditional on log₁₀ total annotation counts. P-value frequencies are nominal, unadjusted for multiplicity and descriptive; they are not combined tests. Percentiles summarize sample-selection sensitivity and are not population confidence intervals. R² values should be compared within each design, because the designs use different sample sizes.')

    page(doc,'Sampling sensitivity and dispersion')
    for x in s['Sampling and annotation-total sensitivity']:paragraph(doc,x)
    for x in s['Dispersion and study structure']:paragraph(doc,x)
    picture(doc,plots/'Supplementary_functional_regional_diagnostics.png',7)
    paragraph(doc,'Supplementary figure Annotation totals and dispersion diagnostics',True)
    paragraph(doc,'(a) Unadjusted and conditional unique realm R² for each prevalence class across the 499 primary selections. The dashed line is equality. (b) PERMDISP permutation P-values for the same selections; the dashed line marks 0.05. Boxes show medians and interquartile ranges; whiskers extend to observations within 1.5 interquartile ranges. Outlying points are omitted for legibility. These diagnostics do not remove confounding by study.')

    page(doc,'Balanced designs and paired Shell sensitivities')
    picture(doc,plots/'Supplementary_functional_regional_sensitivity.png',6.7)
    paragraph(doc,'Supplementary figure Balanced designs and paired Shell contrasts',True)
    paragraph(doc,'(a) Regional R² distributions for Prevalent, Shell and Cloud under primary three-realm × six-locality and alternative three-realm × four-locality designs (499 selections each). Boxes indicate medians and interquartile ranges; whiskers extend to observations within 1.5 interquartile ranges. (b) Shell R² minus the larger of Prevalent and Cloud R², calculated within each selection. Filled points and lines show medians and 5th–95th percentiles for the two designs and annotation-total adjustment (499 selections), and feature-count, occurrence and prevalence sensitivities (99). Open diamonds and intervals show medians and full ranges across six eligible country-deletion medians and 42 study/BioProject-deletion medians. Each deletion used 99 selections. Positive contrasts favour Shell; these intervals describe sensitivity rather than population confidence.')

    page(doc,'Selected Shell process scores')
    for x in s['Selected Shell process scores']:paragraph(doc,x)
    page(doc,'Process effect sizes and catalog')
    effect=pd.read_csv(tabs/'181_process_regional_summary.tsv',sep='\t')
    z=effect[effect.design=='3 x 6']
    table(doc,['Selected process (KOs)','Regional R²\nmedian (5th–95th)','Conditional unique R²\nmedian (5th–95th)'],[[r.process_label+f' ({r.n_KOs})',f'{r.R2_median:.3f} ({r.R2_q05:.3f}–{r.R2_q95:.3f})',f'{r.adjusted_unique_R2_median:.3f} ({r.adjusted_unique_R2_q05:.3f}–{r.adjusted_unique_R2_q95:.3f})'] for r in z.itertuples()],[2.6,2.2,2.2])
    paragraph(doc,'Each row summarizes the same 499 primary selections. Conditional R² accounts for log₁₀ total KO counts. Values measure the regional association of each score, not its contribution to total Shell compositional variation. No process-level significance test was performed.')
    paragraph(doc,'KO membership retained in the current Shell scores',True)
    cat=pd.read_csv(tabs/'181_process_catalog_validation.tsv',sep='\t')
    rows=[]
    for label,g in cat[cat.included_in_score].groupby('process_label',sort=False):
        rows.append([label,', '.join(g.KO)])
    table(doc,['Selected process','Shell KOs'],rows,[2.6,4.4])
    paragraph(doc,'K00441 was excluded from methanogenic/Hdr electron transfer because its current prevalence is 95.29%. The original catalog contained 34 KOs; the current Shell-only catalog contains 33. The complete inclusion audit and alternative-design effect sizes accompany the analysis tables.')

    page(doc,'Locality Shell turnover')
    for x in s['Locality Shell turnover']:paragraph(doc,x)

    page(doc,'Locality turnover and resampling uncertainty')
    picture(doc,plots/'Supplementary_shell_locality_turnover.png',7)
    paragraph(doc,'Supplementary figure Shell turnover across all named localities',True)
    paragraph(doc,'Points show direct RMS Aitchison distance from each locality mean CLR profile to the equally weighted global centroid of all 26 localities. Lines show 5th–95th percentiles from 500 within-locality profile bootstraps that recompute the global centroid. n indicates profiles. Colours distinguish Central Indo-Pacific (blue), Tropical Atlantic (orange), Western Indo-Pacific (green) and Tropical Eastern Pacific (purple). Open diamonds and asterisks identify localities without coordinates, retained in calculations but omitted from the geographic map. Percentiles describe within-locality resampling variation and can exclude the direct estimator; they do not include study or locality selection uncertainty.')
    doc.save(a.outdir/'Functional_regional_supplementary_results.docx')
    print('Replaced both Word reports: 11-point black text; bold titles only; no headers, footers, shading or inline emphasis.')


if __name__=='__main__':main()
