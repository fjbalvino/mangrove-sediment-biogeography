#!/usr/bin/env python3
"""Build the separate Figure 5 Word reports from the canonical manuscript text."""
import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def document():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name != 'Normal'
        style.font.italic = False
        style.paragraph_format.line_spacing = 1
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.space_before = Pt(8 if name.startswith('Heading') else 0)
        style.paragraph_format.keep_with_next = name != 'Normal'
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.core_properties.author = ''
    return doc


def table(doc, rows):
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', c.strip()) for c in row)]
    count = len(rows[0])
    if any(len(row) != count for row in rows):
        raise ValueError('Inconsistent manuscript table column count')
    # Numeric result tables remain plain, with 11-point text throughout.
    weights = [min(25, max(len(c) for c in col)) for col in zip(*rows)]
    widths = [7 * (0.6 + w / 20) / sum(0.6 + v / 20 for v in weights) for w in weights]
    t = doc.add_table(rows=0, cols=count)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        element = OxmlElement('w:' + side)
        for attr, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            element.set(qn('w:' + attr), value)
        borders.append(element)
    t._tbl.tblPr.append(borders)
    for i, values in enumerate(rows):
        row = t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = i < len(rows) - 1 if len(rows) <= 8 else i == 0
            run = p.add_run(value)
            run.font.name, run.font.size = 'Arial', Pt(11)
            run.font.bold = False
            run.font.color.rgb = RGBColor(0, 0, 0)
    gap = doc.add_paragraph()
    gap.paragraph_format.line_spacing = Pt(1)
    gap.paragraph_format.space_after = Pt(4)
    gap.add_run().font.size = Pt(1)


def convert(source, target, figure=None):
    doc = document()
    lines = source.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or re.fullmatch(r'\[.*\]\(.*\)', line):
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            table(doc, rows)
            continue
        if line.startswith('#'):
            depth = len(line) - len(line.lstrip('#'))
            text = line.lstrip('#').strip()
            style = 'Title' if depth == 1 else 'Heading ' + str(min(3, depth - 1))
            doc.add_paragraph(text, style)
        elif line.startswith('**') and line.endswith('**'):
            doc.add_paragraph(line[2:-2], 'Title')
        else:
            if figure and line.startswith('Figure 5 |'):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(figure), width=Inches(6.4))
            doc.add_paragraph(line)
        i += 1
    doc.save(target)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--repository', type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument('--figure', type=Path, required=True)
    parser.add_argument('--outdir', type=Path, required=True)
    args = parser.parse_args()
    if not args.figure.is_file():
        raise FileNotFoundError(args.figure)
    args.outdir.mkdir(parents=True, exist_ok=True)
    for stem, figure in [('FUNCTIONAL_NETWORK_RESULTS', args.figure),
                         ('FUNCTIONAL_NETWORK_SUPPLEMENTARY_RESULTS', None)]:
        output = args.outdir / (stem.capitalize() + '.docx')
        convert(args.repository / 'docs' / (stem + '.md'), output, figure)
        print(output)


if __name__ == '__main__':
    main()
